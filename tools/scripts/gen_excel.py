"""FF14 EXDViewer 数据表内容Diff总结表 - Excel 生成工具（openpyxl）。
用法: gen_excel_tool.exe <json> <out.xlsx>
"""
import json, sys, os
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

def main():
    data = json.load(open(sys.argv[1], encoding='utf-8'))
    out = sys.argv[2]
    wb = Workbook(); ws = wb.active; ws.title = data['sheet_name'][:31]
    cols = data['columns']
    ws.append(['Row'] + [c['name'] for c in cols])
    for i, c in enumerate(cols, start=2):
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(60, c['width'] / 4))
    ws.column_dimensions['A'].width = 10
    for row in data['rows']:
        rec = [row['row']]
        for cell in row['cells']:
            if isinstance(cell, dict) and 'value' in cell:
                rec.append(cell['value'])
            else:
                rec.append('')
        ws.append(rec)
    # 图片列：叠加缩略图
    for r, row in enumerate(data['rows'], start=2):
        for ci, cell in enumerate(row['cells']):
            if isinstance(cell, dict) and 'icon_path' in cell and os.path.exists(cell['icon_path']):
                col_letter = get_column_letter(ci + 2)
                try:
                    img = XLImage(cell['icon_path'])
                    img.width = 30; img.height = 30
                    ws.add_image(img, f"{col_letter}{r}")
                    ws[f"{col_letter}{r}"] = cell['id']
                except Exception:
                    ws[f"{col_letter}{r}"] = cell['id']
    wb.save(out)
    print("saved", out)

main()
